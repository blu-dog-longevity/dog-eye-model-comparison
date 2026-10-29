import csv
import html
import json
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw


ROOT = Path(__file__).parent
REPORT_DIR = ROOT / "runs/report"
DETECTION_REVIEW = {
    ("02_dalmatian_eye.jpg", 2): "Manual review: duplicate detection of the same visible eye.",
    ("04_cataract_dog.jpg", 2): "Manual review: false positive on the nose, not an eye.",
}


PAGE_HEAD = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Canine eye model outputs</title>
<style>
:root { color-scheme: light; --ink:#172026; --muted:#607078; --line:#dce5e7; --paper:#fff; --wash:#f2f6f6; --teal:#007f73; --gold:#d89c00; }
* { box-sizing:border-box; }
body { margin:0; font:15px/1.45 system-ui,-apple-system,sans-serif; color:var(--ink); background:var(--wash); }
header { padding:28px clamp(18px,4vw,56px) 18px; background:#102d31; color:white; }
header h1 { margin:0 0 8px; font-size:clamp(24px,4vw,38px); }
header p { margin:4px 0; max-width:1000px; color:#d8e9e8; }
header a { color:#a8eee4; }
.notice { margin-top:14px; padding:10px 12px; border-left:4px solid var(--gold); background:#183b3f; }
.tabs { position:sticky; top:0; z-index:5; display:flex; gap:8px; overflow:auto; padding:12px clamp(18px,4vw,56px); background:white; border-bottom:1px solid var(--line); }
.tabs button { border:1px solid var(--line); border-radius:999px; padding:9px 14px; background:white; color:var(--ink); font-weight:700; cursor:pointer; white-space:nowrap; }
.tabs button.active { color:white; border-color:var(--teal); background:var(--teal); }
.tabs small { opacity:.75; }
main { padding:22px clamp(18px,4vw,56px) 60px; }
.panel { display:none; }
.panel.active { display:block; }
.panel > h2 { margin:0 0 4px; }
.panel > p { margin:0 0 18px; color:var(--muted); }
.cards { display:grid; grid-template-columns:repeat(auto-fit,minmax(min(100%,520px),1fr)); gap:18px; }
#targeted .cards { grid-template-columns:1fr; }
.card { overflow:hidden; border:1px solid var(--line); border-radius:14px; background:var(--paper); box-shadow:0 3px 12px #14343a10; }
.card-head { display:flex; justify-content:space-between; gap:14px; align-items:start; padding:14px 16px; border-bottom:1px solid var(--line); }
.card h3 { margin:0; font-size:16px; overflow-wrap:anywhere; }
.meta { margin-top:3px; color:var(--muted); font-size:13px; }
.pill { flex:none; border-radius:999px; padding:5px 9px; color:white; background:var(--teal); font-size:12px; font-weight:800; }
.pill.miss { background:#8b3c2f; }
.pill.fail { background:#7c4b00; }
.review { margin-top:8px; padding:8px 10px; border-left:3px solid var(--gold); color:#6b4a00; background:#fff8df; font-size:13px; }
.media { display:grid; grid-template-columns:repeat(auto-fit,minmax(210px,1fr)); gap:1px; background:var(--line); }
figure { margin:0; padding:10px; background:white; }
figcaption { margin:0 0 8px; color:var(--muted); font-size:12px; font-weight:800; text-transform:uppercase; letter-spacing:.04em; }
img { display:block; width:100%; height:330px; object-fit:contain; background:#eef2f2; border-radius:8px; }
.empty { display:grid; place-items:center; height:330px; padding:20px; color:var(--muted); text-align:center; background:#eef2f2; border-radius:8px; }
.details { padding:14px 16px 16px; }
.scores { display:grid; gap:7px; }
.score-row { display:grid; grid-template-columns:minmax(115px,1fr) 3fr 58px; gap:8px; align-items:center; font-size:13px; }
.bar { height:8px; overflow:hidden; border-radius:999px; background:#e8eeee; }
.bar span { display:block; height:100%; background:var(--teal); }
details summary { cursor:pointer; font-weight:700; margin-bottom:10px; }
.metrics { display:grid; grid-template-columns:repeat(auto-fit,minmax(140px,1fr)); gap:9px; }
.metric { padding:9px 10px; border-radius:8px; background:#f4f7f7; }
.metric b { display:block; margin-top:2px; font-size:16px; }
.download { display:inline-block; margin-top:18px; color:var(--teal); font-weight:800; }
@media (max-width:600px) { img,.empty { height:250px; } .card-head { display:block; } .pill { display:inline-block; margin-top:8px; } }
</style>
</head>
<body>
<header>
  <h1>Canine eye model outputs</h1>
  <p><b>Question:</b> Can one compact report make every model's actual behavior visually inspectable?</p>
  <p><b>Audience:</b> Logan and the BLU research team.</p>
  <p>August 26, 2026 archive · <a href="../../../sources.html">Sources &amp; provenance</a> · <a href="../../../SOURCES.md" download>Markdown source record</a> · <a href="https://github.com/blu-dog-longevity/dog-eye-model-comparison">GitHub repository</a></p>
  <div class="notice">Prototype only. The source images are not veterinary ground truth, and no prediction or visual metric is clinically validated.</div>
</header>
"""


PAGE_TAIL = """
<script>
const buttons = document.querySelectorAll('[data-tab]');
const panels = document.querySelectorAll('.panel');
buttons.forEach(button => button.addEventListener('click', () => {
  buttons.forEach(item => item.classList.toggle('active', item === button));
  panels.forEach(panel => panel.classList.toggle('active', panel.id === button.dataset.tab));
}));
</script>
</body>
</html>
"""


def web_path(path: Path) -> str:
    return os.path.relpath(path, REPORT_DIR).replace(os.sep, "/")


def image_figure(label: str, path: Path, alt: str) -> str:
    return (
        f"<figure><figcaption>{html.escape(label)}</figcaption>"
        f'<img loading="lazy" src="{html.escape(web_path(path))}" '
        f'alt="{html.escape(alt)}"></figure>'
    )


def score_rows(scores: dict[str, float]) -> str:
    rows = []
    for name, score in sorted(scores.items(), key=lambda item: item[1], reverse=True):
        percent = score * 100
        rows.append(
            '<div class="score-row">'
            f"<span>{html.escape(name)}</span>"
            f'<span class="bar"><span style="width:{percent:.2f}%"></span></span>'
            f"<b>{percent:.1f}%</b></div>"
        )
    return "".join(rows)


def boxed_image(row: dict) -> Path:
    output = REPORT_DIR / "assets/boxes" / (
        f"{Path(row['image']).stem}_eye_{row['eye']}.jpg"
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    image = Image.open(ROOT / "data/images" / row["image"]).convert("RGB")
    draw = ImageDraw.Draw(image)
    draw.rectangle(row["box"], outline=(255, 196, 0), width=max(2, image.width // 500))
    image.save(output, quality=88)
    return output


def rady_card(row: dict) -> str:
    status_ok = row["status"] == "ok"
    eye_label = f"Eye {row['eye']}" if status_ok else "No eye crop"
    pill_class = "pill" if status_ok else "pill miss"
    pill_text = row["prediction"] if status_ok else "No eye detected"
    original = ROOT / "data/images" / row["image"]
    media = image_figure(
        "Before crop",
        boxed_image(row) if status_ok else original,
        f"Original {row['image']}",
    )
    if status_ok:
        crop = ROOT / "runs/first-pass" / row["crop"]
        media += image_figure("After crop", crop, f"Eye crop from {row['image']}")
        details = f'<div class="scores">{score_rows(row["scores"])}</div>'
        meta = f"Detection {row['detection_score'] * 100:.1f}% · Classification {row['score'] * 100:.1f}%"
    else:
        media += '<figure><figcaption>After crop</figcaption><div class="empty">YOLO produced no crop for this image.</div></figure>'
        details = ""
        meta = "Classifier did not run because the detector produced no crop."
    review = DETECTION_REVIEW.get((row["image"], row.get("eye")))
    review_html = f'<div class="review">{html.escape(review)}</div>' if review else ""
    return (
        '<article class="card">'
        f'<div class="card-head"><div><h3>{html.escape(row["image"])} — {eye_label}</h3>'
        f'<div class="meta">{meta}</div>{review_html}</div><span class="{pill_class}">{html.escape(pill_text)}</span></div>'
        f'<div class="media">{media}</div><div class="details">{details}</div></article>'
    )


def dog_card(row: dict) -> str:
    original = ROOT / "data/images" / row["image"]
    return (
        '<article class="card">'
        f'<div class="card-head"><div><h3>{html.escape(row["image"])}</h3>'
        '<div class="meta">Whole image; no eye detector or crop.</div></div>'
        f'<span class="pill">{html.escape(row["prediction"])}</span></div>'
        f'<div class="media">{image_figure("Whole-image input", original, row["image"])}</div>'
        '<div class="details"><details><summary>All 20 output scores</summary>'
        f'<div class="scores">{score_rows(row["scores"])}</div></details></div></article>'
    )


def metric(label: str, value: str) -> str:
    return f'<div class="metric"><span>{html.escape(label)}</span><b>{html.escape(value)}</b></div>'


def targeted_card(row: dict) -> str:
    success = row["status"] == "candidate"
    pill_class = "pill" if success else "pill fail"
    pill_text = "Provisional masks" if success else "Targeting failed"
    media = "".join(
        [
            image_figure("Manual prompts", ROOT / row["prompt_view"], f"Prompts for {row['id']}"),
            image_figure("Combined overlay", ROOT / row["overlay"], f"Targeted masks for {row['id']}"),
            image_figure("Mask only", ROOT / row["mask_view"], f"Pupil and iris masks for {row['id']}"),
        ]
    )
    if success:
        values = "".join(
            [
                metric("Pupil / iris-disk area", f"{row['pupil_iris_area_ratio'] * 100:.1f}%"),
                metric("Pupil / iris diameter", f"{row['pupil_iris_equivalent_diameter_ratio']:.3f}"),
                metric("Pupil area", f"{row['pupil_area_px']:,} px"),
                metric("Iris disk area", f"{row['iris_disk_area_px']:,} px"),
                metric("Pupil asymmetry", f"{row['pupil_shape_asymmetry']:.3f}"),
                metric("Pupil center offset", f"{row['pupil_center_offset_iris_radius']:.3f} iris radii"),
                metric("Iris mean RGB", f"{row['iris_mean_red']:.0f}, {row['iris_mean_green']:.0f}, {row['iris_mean_blue']:.0f}"),
                metric("Iris saturation", f"{row['iris_mean_saturation']:.3f}"),
                metric("Pupil SAM score", f"{row['pupil_prompt']['selected_score']:.3f}"),
                metric("Iris SAM score", f"{row['iris_prompt']['selected_score']:.3f}"),
            ]
        )
        details = f'<div class="metrics">{values}</div><div class="review">Manual candidate selection; these ratios are not ground truth.</div>'
    else:
        details = (
            f'<div class="review">{html.escape(row["reason"])} No pupil/iris metric was calculated.</div>'
        )
    return (
        '<article class="card">'
        f'<div class="card-head"><div><h3>{html.escape(row["image"])} — {html.escape(row["eye"])}</h3>'
        '<div class="meta">Magenta: pupil attempt · cyan: iris attempt · red points: exclusions</div></div>'
        f'<span class="{pill_class}">{pill_text}</span></div>'
        f'<div class="media">{media}</div><div class="details">{details}</div></article>'
    )


def flatten_rady(rows: list[dict], model: str) -> list[dict]:
    output = []
    for row in rows:
        if row["model"] != model:
            continue
        flat = {
            "image": row["image"],
            "eye": row.get("eye", ""),
            "status": row["status"],
            "prediction": row.get("prediction", ""),
            "score": row.get("score", ""),
            "detection_score": row.get("detection_score", ""),
            "original_path": f"data/images/{row['image']}",
            "crop_path": f"runs/first-pass/{row.get('crop', '')}" if row.get("crop") else "",
        }
        box = row.get("box", ["", "", "", ""])
        flat.update(dict(zip(["box_x1", "box_y1", "box_x2", "box_y2"], box)))
        flat.update(row.get("scores", {}))
        output.append(flat)
    return output


def flatten_dog(rows: list[dict]) -> list[dict]:
    output = []
    for row in rows:
        if row["model"] != "dog_disease_predictor":
            continue
        flat = {
            "image": row["image"],
            "status": row["status"],
            "prediction": row["prediction"],
            "score": row["score"],
            "original_path": f"data/images/{row['image']}",
        }
        flat.update(row["scores"])
        output.append(flat)
    return output


def flatten_sam(rows: list[dict]) -> list[dict]:
    output = []
    for row in rows:
        flat = {key: value for key, value in row.items() if key != "box"}
        flat.update(dict(zip(["box_x1", "box_y1", "box_x2", "box_y2"], row["box"])))
        flat["original_path"] = f"data/images/{row['image']}"
        output.append(flat)
    return output


def flatten_targeted(rows: list[dict]) -> list[dict]:
    metric_names = [
        "pupil_area_px",
        "iris_disk_area_px",
        "pupil_iris_area_ratio",
        "pupil_iris_equivalent_diameter_ratio",
        "pupil_shape_asymmetry",
        "pupil_center_offset_iris_radius",
        "iris_mean_red",
        "iris_mean_green",
        "iris_mean_blue",
        "iris_mean_saturation",
    ]
    output = []
    for row in rows:
        flat = {
            "id": row["id"],
            "image": row["image"],
            "eye": row["eye"],
            "status": row["status"],
            "reason": row["reason"],
            "pupil_selected_candidate": row["pupil_prompt"]["selected_candidate"],
            "pupil_sam_score": row["pupil_prompt"]["selected_score"],
            "iris_selected_candidate": row["iris_prompt"]["selected_candidate"],
            "iris_sam_score": row["iris_prompt"]["selected_score"],
            "prompt_view": row["prompt_view"],
            "overlay": row["overlay"],
            "mask_view": row["mask_view"],
        }
        flat.update({name: row.get(name, "") for name in metric_names})
        output.append(flat)
    return output


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def write_workbook(sheets: dict[str, list[dict]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        headers = list(rows[0])
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="007F73")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(48, max(len(str(cell.value or "")) for cell in column) + 2)
            sheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(REPORT_DIR / "model_outputs.xlsx")


def main() -> None:
    results = json.loads((ROOT / "runs/first-pass/results.json").read_text())
    sam_rows = json.loads((ROOT / "runs/sam2/metrics.json").read_text())
    targeted_rows = json.loads((ROOT / "runs/targeted-sam/results.json").read_text())
    REPORT_DIR.mkdir(parents=True, exist_ok=True)

    rady4 = [row for row in results if row["model"] == "rady4"]
    rady6 = [row for row in results if row["model"] == "rady6"]
    dog = [row for row in results if row["model"] == "dog_disease_predictor"]
    panels = {
        "rady4": ("Rady-4", rady4, rady_card, "YOLO eye detection followed by the four-label classifier."),
        "rady6": ("Rady-6", rady6, rady_card, "The same YOLO checkpoint followed by the separate six-label classifier."),
        "dog": ("DogDiseasePredictor", dog, dog_card, "Whole-image 20-label predictions; no eye crop."),
        "targeted": ("Targeted SAM", targeted_rows, targeted_card, "Manual boxes and positive/negative clicks for pupil and iris. This tests feasibility; it is not automated or validated."),
    }

    tabs = ['<nav class="tabs">']
    sections = ["<main>"]
    for index, (key, (title, rows, renderer, description)) in enumerate(panels.items()):
        active = " active" if index == 0 else ""
        tabs.append(
            f'<button class="{active.strip()}" data-tab="{key}">{html.escape(title)} <small>({len(rows)})</small></button>'
        )
        cards = "".join(renderer(row) for row in rows)
        sections.append(
            f'<section id="{key}" class="panel{active}"><h2>{html.escape(title)}</h2>'
            f'<p>{html.escape(description)}</p><div class="cards">{cards}</div></section>'
        )
    tabs.append("</nav>")
    sections.append('<a class="download" href="model_outputs.xlsx">Download the Excel workbook</a>')
    sources = json.loads((ROOT / "sources.json").read_text())["images"]
    sections.append('<h2 id="image-credits">Image credits</h2><p>Crops, boxes, masks, and overlays are BLU Dog annotations of these photographs. CC BY-SA adaptations retain their respective source licenses. <a href="../../../sources.html">Full provenance and licenses</a>.</p><ul>')
    for source in sources:
        sections.append(
            f'<li><a href="{html.escape(source["page"])}">{html.escape(source["file"])}</a> — '
            f'{html.escape(source["creator"])} · {html.escape(source["license"])}</li>'
        )
    sections.append('</ul></main>')
    (REPORT_DIR / "index.html").write_text(
        PAGE_HEAD + "".join(tabs) + "".join(sections) + PAGE_TAIL
    )

    data_dir = REPORT_DIR / "csv"
    data_dir.mkdir(exist_ok=True)
    sheets = {
        "Rady-4": flatten_rady(results, "rady4"),
        "Rady-6": flatten_rady(results, "rady6"),
        "DogDisease": flatten_dog(results),
        "SAM2-Metrics": flatten_sam(sam_rows),
        "Targeted-SAM": flatten_targeted(targeted_rows),
    }
    for name, rows in sheets.items():
        write_csv(data_dir / f"{name.lower().replace('-', '_')}.csv", rows)
    write_workbook(sheets)
    print(REPORT_DIR / "index.html")
    print(REPORT_DIR / "model_outputs.xlsx")


if __name__ == "__main__":
    main()
